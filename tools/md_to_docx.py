#!/usr/bin/env python3
"""Minimal Markdown -> .docx converter for the report files (python-docx). Supports #/##/### headings, paragraphs with
**bold**, *italic* and `code`, pipe tables, bullet (*, -) and numbered (1.) lists, and images (![](path)).
Lines inside one blank-line block become separate paragraphs. Usage: md_to_docx.py in.md out.docx [--font Calibri --size 11]"""
import argparse, os, re, sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

INLINE = re.compile(r"(\*\*\*.+?\*\*\*|\*\*.+?\*\*|\*(?!\s)[^*\n]+?\*|`[^`\n]+`)")

def add_runs(par, text, base_bold=False, base_italic=False, size=None):
    for piece in INLINE.split(text):
        if not piece: continue
        bold, italic, code = base_bold, base_italic, False
        if piece.startswith("***") and piece.endswith("***") and len(piece) > 6: piece, bold, italic = piece[3:-3], True, True
        elif piece.startswith("**") and piece.endswith("**") and len(piece) > 4: piece, bold = piece[2:-2], True
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2: piece, italic = piece[1:-1], True
        elif piece.startswith("`") and piece.endswith("`") and len(piece) > 2: piece, code = piece[1:-1], True
        run = par.add_run(piece); run.bold = bold or None; run.italic = italic or None
        if code:
            run.font.name = "Consolas"; run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            if size: run.font.size = Pt(size - 1)
        elif size: run.font.size = Pt(size)

def add_table(doc, lines, size):
    rows = [l.strip() for l in lines if l.strip().startswith("|")]
    rows = [r for r in rows if not re.match(r"^\|[\s:|-]+\|$", r)]
    cells = [[c.strip() for c in r.strip("|").split("|")] for r in rows]
    ncol = max(len(r) for r in cells)
    t = doc.add_table(rows=len(cells), cols=ncol); t.style = "Table Grid"
    for i, r in enumerate(cells):
        for j in range(ncol):
            txt = r[j] if j < len(r) else ""
            cell = t.cell(i, j); cell.text = ""; p = cell.paragraphs[0]
            add_runs(p, txt, base_bold=(i == 0), size=size - 2)
    doc.add_paragraph()

def convert(src, dst, font="Calibri", size=11, imgwidth=6.3):
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = font; st.font.size = Pt(size); st.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    for sec in doc.sections: sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Inches(1)
    base = os.path.dirname(os.path.abspath(src)); text = open(src, encoding="utf-8").read()
    for block in re.split(r"\n\s*\n", text):
        block = block.strip("\n")
        if not block.strip(): continue
        lines = block.split("\n")
        if all(l.strip().startswith("|") for l in lines if l.strip()):
            add_table(doc, lines, size); continue
        for line in lines:
            s = line.rstrip()
            if not s.strip(): continue
            m = re.match(r"^(#{1,6})\s+(.*)$", s)
            if m:
                level = len(m.group(1)); h = doc.add_heading("", level=0 if level == 1 else min(level - 1, 3)); add_runs(h, m.group(2)); continue
            m = re.match(r"^!\[[^\]]*\]\(([^)]+)\)\s*$", s)
            if m:
                path = m.group(1) if os.path.isabs(m.group(1)) else os.path.join(base, m.group(1))
                if os.path.exists(path): doc.add_picture(path, width=Inches(imgwidth)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                else: doc.add_paragraph(f"[missing image: {m.group(1)}]")
                continue
            m = re.match(r"^\s*[*-]\s+(.*)$", s)
            if m: add_runs(doc.add_paragraph(style="List Bullet"), m.group(1), size=size); continue
            m = re.match(r"^\s*\d+\.\s+(.*)$", s)
            if m: add_runs(doc.add_paragraph(style="List Number"), m.group(1), size=size); continue
            p = doc.add_paragraph(); add_runs(p, s, size=size); p.paragraph_format.space_after = Pt(3)
    doc.save(dst); print("saved", dst)

if __name__ == "__main__":
    ap = argparse.ArgumentParser(); ap.add_argument("src"); ap.add_argument("dst"); ap.add_argument("--font", default="Calibri"); ap.add_argument("--size", type=float, default=11); ap.add_argument("--imgwidth", type=float, default=6.3)
    a = ap.parse_args(); convert(a.src, a.dst, a.font, a.size, a.imgwidth)
